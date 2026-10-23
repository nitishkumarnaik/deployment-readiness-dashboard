import random, datetime as dt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

random.seed(7)
today = dt.date(2026, 10, 23)

sites = [("S01", "Limerick", "Ireland"), ("S02", "Athlone", "Ireland"), ("S03", "Dundalk", "Ireland"),
         ("S04", "Austin", "USA"), ("S05", "Raleigh", "USA"), ("S06", "Singapore", "Singapore"),
         ("S07", "Eindhoven", "Netherlands"), ("S08", "Heredia", "Costa Rica")]
golive, rows = {}, []
start = dt.date(2026, 9, 7)
for i, (sid, name, reg) in enumerate(sites):
    g = start + dt.timedelta(weeks=i * 4)
    golive[sid] = g
    rows.append(dict(SiteID=sid, SiteName=name, Region=reg, Wave=f"Wave {i // 3 + 1}",
                     PlannedGoLive=g, DeploymentLead=random.choice(["A. Murphy", "J. Walsh", "P. Ryan", "S. Kelly"])))
df_sites = pd.DataFrame(rows)

streams = {"Training": ["Super-user training", "End-user training", "Training sign-off"],
           "Process": ["Current-state process mapped", "Gap analysis vs standard", "Future-state SOPs approved"],
           "Systems": ["User access set up", "Integration tested", "Cutover plan approved"],
           "Data Migration": ["Data extract", "Data cleansing", "Mock load", "Final load validated"],
           "Validation": ["IQ/OQ complete", "PQ complete", "Validation report signed"]}

perf = {"S03": 0.55, "S04": 0.8, "S05": 0.7}  # sites that are behind plan
tasks, n = [], 1
for sid, _, _ in sites:
    g = golive[sid]
    for ws, items in streams.items():
        for j, t in enumerate(items):
            planned = g - dt.timedelta(days=84 - j * 18 - random.randint(0, 10))
            if planned <= today and random.random() < perf.get(sid, 0.92):
                comp = min(planned + dt.timedelta(days=random.randint(-5, 14)), today)
                status = "Complete"
            else:
                comp = None
                status = "Overdue" if planned < today else random.choice(["Not started", "In progress"])
            tasks.append(dict(TaskID=f"T{n:04d}", SiteID=sid, Workstream=ws, Task=t, PlannedDate=planned,
                              CompletedDate=comp, Status=status))
            n += 1
df_tasks = pd.DataFrame(tasks)

descs = {"Training": "Super-users not released from production", "Process": "Local SOP differs from global standard",
         "Systems": "Integration test failures", "Data Migration": "Master data quality below threshold",
         "Validation": "Validation resource shortage"}
risks, n = [], 1
for sid, _, _ in sites:
    for _ in range(random.randint(2, 6)):
        ws = random.choice(list(streams))
        raised = today - dt.timedelta(days=random.randint(5, 120))
        closed = raised + dt.timedelta(days=random.randint(5, 40)) if random.random() < 0.4 else None
        if closed and closed > today:
            closed = None
        risks.append(dict(RiskID=f"R{n:03d}", SiteID=sid, Workstream=ws, Description=descs[ws],
                          Severity=random.choice(["High", "High", "Medium", "Medium", "Low"]),
                          Status="Closed" if closed else "Open", RaisedDate=raised, ClosedDate=closed))
        n += 1
df_risks = pd.DataFrame(risks)

ai = []
for k in range(40):
    d = today - dt.timedelta(days=random.randint(1, 90))
    errs = random.choice([0, 0, 0, 1, 1, 2, 3])
    ai.append(dict(ReportID=f"AI{k + 1:03d}", ReportDate=d, SiteID=random.choice(sites)[0],
                   ReportType=random.choice(["Readiness summary", "Process comparison", "Risk summary"]),
                   ClaimsChecked=random.randint(15, 40), ErrorsFound=errs,
                   MainErrorType="None" if errs == 0 else random.choice(["Wrong figure", "Hallucinated item", "Outdated data"])))
df_ai = pd.DataFrame(ai).sort_values("ReportDate")

readme = pd.DataFrame({"Item": ["What this is", "Sheets", "Relationships", "Snapshot date"],
                       "Detail": ["Synthetic sample data for a Power BI deployment readiness dashboard (portfolio project). All names and figures are made up.",
                                  "Sites, ReadinessTasks, Risks, AI_Report_Checks",
                                  "Sites[SiteID] one-to-many with ReadinessTasks[SiteID], Risks[SiteID] and AI_Report_Checks[SiteID]",
                                  "'Today' in this data is 23 October 2026"]})

out = "data/Deployment_Readiness_Sample_Data.xlsx"
with pd.ExcelWriter(out, engine="openpyxl", date_format="YYYY-MM-DD", datetime_format="YYYY-MM-DD") as w:
    readme.to_excel(w, sheet_name="README", index=False)
    df_sites.to_excel(w, sheet_name="Sites", index=False)
    df_tasks.to_excel(w, sheet_name="ReadinessTasks", index=False)
    df_risks.to_excel(w, sheet_name="Risks", index=False)
    df_ai.to_excel(w, sheet_name="AI_Report_Checks", index=False)

wb = load_workbook(out)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10, bold=(c.row == 1))
            if hasattr(c.value, "year"):
                c.number_format = "YYYY-MM-DD"
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    for col in ws.columns:
        L = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, L + 2), 90)
    ws.freeze_panes = "A2"
wb.properties.creator = "S M Nitish Kumar Naik"
wb.properties.lastModifiedBy = "S M Nitish Kumar Naik"
wb.save(out)

done = df_tasks.groupby("SiteID").Status.apply(lambda s: round((s == "Complete").mean() * 100))
print(len(df_sites), len(df_tasks), len(df_risks), len(df_ai))
print(df_tasks.Status.value_counts().to_dict())
print(done.to_dict())
